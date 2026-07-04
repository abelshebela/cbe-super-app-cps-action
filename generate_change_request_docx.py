from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re


def set_cell_shading(cell, fill):
    """Set background shading for a table cell."""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_hyperlink(paragraph, text, url):
    """Add a hyperlink to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rPr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_heading_custom(doc, text, level=1):
    """Add a heading with custom formatting."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        if level == 1:
            run.font.size = Pt(16)
            run.font.bold = True
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
    return heading


def main():
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run('Change Request — CBE Super App CPS Action')
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    doc.add_paragraph()

    # Intro
    p = doc.add_paragraph('Hi ')
    p.add_run('DevOps / Release Management Team').bold = True
    p.add_run(',')
    doc.add_paragraph()

    p = doc.add_paragraph(
        'We are requesting a review of the Production and Staging environment synchronization for the '
    )
    p.add_run('CBE Super App — CPS Action').bold = True
    p.add_run(
        ' project. Below is the summary of the current repository state and the proposed action for your review and approval.'
    )
    doc.add_paragraph()

    # Overview
    add_heading_custom(doc, '📋 Overview', level=2)
    overview_items = [
        ('Requested By:', 'CBE Super App CPS Backend Team'),
        ('Priority:', '🟢 Low'),
        ('Target Deployment:', 'N/A — no deployment required at this time'),
        ('Source Branches:', 'origin/prod ↔ origin/staging comparison'),
        ('Repository Status:', 'origin/prod contains 8 merge commits not present in origin/staging; origin/staging contains 0 code commits not present in origin/prod; No file-level differences between Production and Staging'),
        ('Associated Ticket(s):', 'N/A'),
    ]
    for label, value in overview_items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(label).bold = True
        p.add_run(' ' + value)
    doc.add_paragraph()

    # Description
    add_heading_custom(doc, '🔍 Description & Justification', level=2)
    p = doc.add_paragraph(
        'A branch comparison between '
    )
    p.add_run('origin/prod').bold = True
    p.add_run(' and ')
    p.add_run('origin/staging').bold = True
    p.add_run(
        ' shows that the Production and Staging environments are currently '
    )
    p.add_run('code-identical').bold = True
    p.add_run(
        '. The 8 additional commits in Production are merge commits created when Staging was promoted to Production (e.g., '
    )
    p.add_run('Merge pull request #3559 / #3535 / #3495 / #3471 / #3467 / #3459 / #3455 / #3450 from CBE-Super-App/staging').italic = True
    p.add_run(
        '). These merge commits do not introduce any new file changes compared to Staging.'
    )
    doc.add_paragraph()
    p = doc.add_paragraph(
        'No application code, database schema, or configuration changes are pending between Production and Staging. Therefore, no deployment or promotion is required at this time.'
    )
    doc.add_paragraph()

    # Impact
    add_heading_custom(doc, '⚡ Impact & Components Affected', level=2)
    impact_items = [
        ('Code Repository:', 'cbe-super-app-cps-action — Production and Staging are synchronized at the code level'),
        ('Database Schema:', 'No — no new columns, tables, or migrations required'),
        ('Environment Variables:', 'No changes required'),
        ('Compliance Risk:', '🟢 Low — no functional changes are being introduced; only merge commit history differs between branches'),
    ]
    for label, value in impact_items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(label).bold = True
        p.add_run(' ' + value)
    doc.add_paragraph()

    # Deployment Plan
    add_heading_custom(doc, '🛠 Deployment & Rollback Plan', level=2)
    p = doc.add_paragraph()
    p.add_run('Deployment Steps:').bold = True
    doc.add_paragraph('No deployment is required. Production and Staging are currently running the same code.', style='List Bullet')
    doc.add_paragraph(
        'If synchronization of branch merge history is desired, merge the latest Production branch into Staging to align the commit history. No application changes will be introduced by this operation.',
        style='List Bullet'
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('Rollback Plan:').bold = True
    doc.add_paragraph('Not applicable — no changes are being deployed.', style='List Bullet')
    doc.add_paragraph()

    # QA
    add_heading_custom(doc, '🧪 QA / Verification Steps', level=2)
    qa_steps = [
        'Confirm that git diff origin/staging..origin/prod returns no changed files.',
        'Verify that both environments are running the same container image tag.',
        'Confirm application health checks pass on both Production and Staging environments.',
    ]
    for step in qa_steps:
        doc.add_paragraph(step, style='List Number')
    doc.add_paragraph()

    # Branch Summary
    add_heading_custom(doc, '📊 Branch Summary', level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr_cells = table.rows[0].cells
    headers = ['Branch', 'Latest Commit', 'Notes']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_shading(hdr_cells[i], 'D9E1F2')
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    rows = [
        ('origin/dev', '4958953d5', 'Contains 7 commits not in UAT (pending next UAT promotion)'),
        ('origin/uat', '91c793f3d', 'Code-identical with Staging'),
        ('origin/staging', '47cdaa5de', 'Code-identical with Production'),
        ('origin/prod', '5b8145e50', 'Contains 8 merge commits from Staging not present in Staging branch'),
    ]
    for branch, commit, notes in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = branch
        row_cells[1].text = commit
        row_cells[2].text = notes

    doc.add_paragraph()

    # Note
    p = doc.add_paragraph()
    p.add_run('Note:').bold = True
    p.add_run(
        ' The next functional changes available for promotion are currently on the '
    )
    p.add_run('dev').bold = True
    p.add_run(' branch (7 commits ahead of UAT). These are outside the scope of the Production / Staging comparison.')
    doc.add_paragraph()

    # Closing
    p = doc.add_paragraph('Please review and confirm that no action is required, or provide guidance if a different deployment scope is intended.')
    doc.add_paragraph()
    p = doc.add_paragraph('Best regards,')
    doc.add_paragraph()

    # Signature lines
    signatures = [
        '[Requested By]',
        '[Checked By]',
        '[Reviewed By]',
        '[Approved By]',
    ]
    for sig in signatures:
        p = doc.add_paragraph()
        p.add_run(sig).bold = True
        p.add_run(' ______________________ — ______________________ — ______________________')
        doc.add_paragraph()

    doc.save('Change_Request_Prod_Staging_Sync.docx')
    print('Created Change_Request_Prod_Staging_Sync.docx')


if __name__ == '__main__':
    main()
